"""Document archive: watcher versioning, deletion, all-file-type acceptance,
safe archive handling, OCR drafts and long Windows paths."""
from __future__ import annotations

import io
import os
import time
import zipfile

import pytest
from sqlalchemy import select

from conftest import audit_rows, make_pdf


def _scan(settings):
    from app.ingestion.watcher import Watcher

    w = Watcher(settings)
    time.sleep(2.1)  # watcher skips files modified in the last 2 seconds (copy in progress)
    return w.scan()


def _doc_by_relpath(relpath):
    from app.db.models import Document
    from app.db.session import session_scope

    with session_scope() as db:
        d = db.execute(select(Document).where(Document.source_relpath == relpath)).scalar_one()
        db.expunge(d)
        return d


def _versions(doc_id):
    from app.db.models import DocumentVersion
    from app.db.session import session_scope

    with session_scope() as db:
        vs = db.execute(select(DocumentVersion).where(DocumentVersion.document_id == doc_id)
                        .order_by(DocumentVersion.version_number)).scalars().all()
        for v in vs:
            db.expunge(v)
        return vs


def test_watched_file_lifecycle_versioning_deletion_and_reappearance(settings, admin, process_jobs):
    iso = settings.iso_booklets_path
    src = make_pdf(iso / "lifecycle-test.pdf", ["ISO 9999:2020\nLifecycle test document\nThe widget flank angle is defined here."])
    s1 = _scan(settings)
    assert s1["new"] >= 1
    process_jobs()
    doc = _doc_by_relpath("lifecycle-test.pdf")
    assert doc.status == "active" and doc.standard_code == "ISO 9999:2020"
    v1 = _versions(doc.id)
    assert len(v1) == 1 and v1[0].is_active and v1[0].ingestion_status == "indexed"
    raw1 = settings.data_root / v1[0].storage_relpath
    raw1_bytes = raw1.read_bytes()

    # --- change the file -> new immutable version, previous superseded
    make_pdf(src, ["ISO 9999:2020\nLifecycle test document revision 2\nThe widget flank angle changed."])
    s2 = _scan(settings)
    assert s2["changed"] == 1
    process_jobs()
    vs = _versions(doc.id)
    assert [v.version_number for v in vs] == [1, 2]
    assert vs[0].state == "superseded" and not vs[0].is_active
    assert vs[1].state == "active" and vs[1].is_active
    assert raw1.read_bytes() == raw1_bytes, "previous raw version must stay unchanged"
    assert (settings.data_root / vs[1].storage_relpath).exists()
    assert (settings.data_root / "documents" / str(doc.id) / "manifest.json").exists()

    # superseded chunks never appear in retrieval
    r = admin.post("/retrieval/search", {"query": "widget flank angle revision"})
    assert r.status_code == 200

    # --- remove the file -> logical deletion, raw versions kept, excluded from retrieval
    os.remove(src)
    s3 = _scan(settings)
    assert s3["missing"] == 1
    doc = _doc_by_relpath("lifecycle-test.pdf")
    assert doc.status == "deleted" and doc.delete_reason.startswith("source_missing")
    vs = _versions(doc.id)
    assert all(not v.is_active for v in vs)
    assert all((settings.data_root / v.storage_relpath).exists() for v in vs)
    assert audit_rows("document.source_missing", target_id=str(doc.id))

    # --- file reappears -> new version, active again
    make_pdf(src, ["ISO 9999:2020\nLifecycle test document revision 3"])
    s4 = _scan(settings)
    assert s4["reappeared"] == 1
    process_jobs()
    doc = _doc_by_relpath("lifecycle-test.pdf")
    assert doc.status == "active"
    assert [v.version_number for v in _versions(doc.id)] == [1, 2, 3]
    os.remove(src)
    _scan(settings)


def test_long_windows_path_is_ingested(settings, process_jobs):
    deep = settings.iso_booklets_path / ("uzun klasor adi " * 6).strip() / ("alt klasor " * 5).strip()
    name = "Cylindrical gears long filename test -- " + "x" * 90 + " -- Anna’s Archive.pdf"
    full = deep / name
    assert len(str(full)) > 260
    from app.core.paths import fs

    os.makedirs(fs(deep), exist_ok=True)
    tmp = make_pdf(settings.data_root / "tmp" / "long-src.pdf", ["ISO 9998:2019\nLong path content"])
    with open(fs(full), "wb") as f:
        f.write(tmp.read_bytes())
    summary = _scan(settings)
    assert not summary["errors"], summary["errors"]
    process_jobs()
    rel = os.path.relpath(str(full), str(settings.iso_booklets_path)).replace("\\", "/")
    assert _doc_by_relpath(rel).status == "active"
    os.remove(fs(full))
    _scan(settings)


def test_ui_delete_is_logical_and_audited(admin, process_jobs):
    r = admin.upload("/documents/upload", "not.txt", "Dişli ölçüm notu: yanak kontrolü".encode("utf-8"), "text/plain")
    assert r.status_code == 201
    doc_id = r.json()["id"]
    process_jobs()
    d = admin.delete(f"/documents/{doc_id}", {"reason": "test silme"})
    assert d.status_code == 200 and d.json()["status"] == "deleted"
    assert d.json()["current_version"]["state"] == "deleted"
    # raw version is still downloadable (history) and every step is audited
    vid = d.json()["current_version"]["id"]
    dl = admin.get(f"/documents/{doc_id}/versions/{vid}/download")
    assert dl.status_code == 200 and "yanak kontrolü".encode("utf-8") in dl.content
    assert audit_rows("document.delete", target_id=doc_id)
    assert audit_rows("document.download", target_id=doc_id)
    assert admin.post(f"/documents/{doc_id}/reindex").status_code == 409


def _docx_bytes() -> bytes:
    import docx

    d = docx.Document()
    d.add_paragraph("Muayene raporu: helis sapması ölçüldü.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _xlsx_bytes() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.append(["Parça", "Modül"])
    wb.active.append(["Pinyon", 2.5])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.parametrize("name,data,mime,expect_status", [
    ("rapor.docx", None, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "indexed"),
    ("tablo.xlsx", None, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "indexed"),
    ("notlar.md", "# Başlık\nDiş kalınlığı kontrolü".encode("utf-8"), "text/markdown", "indexed"),
    ("eski.doc", b"\xd0\xcf\x11\xe0legacy", "application/msword", "stored_only"),
    ("veri.bin", os.urandom(512), "application/octet-stream", "stored_only"),
    ("arsiv.7z", b"7z\xbc\xaf\x27\x1c" + os.urandom(64), "application/x-7z-compressed", "stored_only"),
])
def test_all_file_types_are_accepted_and_archived(admin, process_jobs, name, data, mime, expect_status):
    if name.endswith(".docx"):
        data = _docx_bytes()
    elif name.endswith(".xlsx"):
        data = _xlsx_bytes()
    r = admin.upload("/attachments", name, data, mime)
    assert r.status_code == 201, r.text
    doc_id = r.json()["id"]
    process_jobs()
    doc = admin.get(f"/documents/{doc_id}").json()
    assert doc["current_version"]["ingestion_status"] == expect_status
    assert doc["is_verified_corpus"] is False  # uploads are never verified corpus
    if expect_status == "stored_only":
        assert doc["current_version"]["extraction_summary"]["stored_only_reason"]


def _zip(entries: dict[str, bytes], compression=zipfile.ZIP_DEFLATED) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression) as z:
        for n, b in entries.items():
            z.writestr(n, b)
    return buf.getvalue()


def test_archive_inventory_and_zip_slip_blocked(admin, process_jobs, settings):
    payload = _zip({
        "ok/aciklama.txt": "Arşiv içi metin: temel adım".encode("utf-8"),
        "../../kacis.txt": b"zip slip attempt",
        "C:/Windows/evil.txt": b"absolute",
        "bomb.txt": b"0" * 2_000_000,  # ratio >> ARCHIVE_MAX_COMPRESSION_RATIO
    })
    r = admin.upload("/attachments", "paket.zip", payload, "application/zip")
    doc_id = r.json()["id"]
    process_jobs()
    doc = admin.get(f"/documents/{doc_id}").json()
    vid = doc["current_version"]["id"]
    members = {m["member_path"]: m for m in admin.get(f"/documents/{doc_id}/versions/{vid}/archive-members").json()}
    assert members["ok/aciklama.txt"]["status"] == "extracted"
    assert members["../../kacis.txt"]["status"] == "blocked"
    assert members["C:/Windows/evil.txt"]["status"] == "blocked"
    assert members["bomb.txt"]["status"] == "blocked"
    # nothing escaped the controlled temp folder and temp folders are cleaned up
    assert not (settings.project_root / "kacis.txt").exists()
    assert not (settings.data_root / "kacis.txt").exists()
    assert not [p for p in (settings.data_root / "tmp").glob("extract-*")]
    pages = admin.get(f"/documents/{doc_id}/versions/{vid}/pages").json()
    assert pages and pages[0]["confidence_status"] == "unverified"


def _text_image(text: str) -> bytes:
    from PIL import Image, ImageDraw, ImageFont

    img = Image.new("RGB", (900, 220), "white")
    d = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 44)
    except OSError:
        font = ImageFont.load_default()
    d.text((20, 60), text, fill="black", font=font)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


@pytest.mark.slow
def test_image_ocr_creates_draft_values_and_draft_page(admin, process_jobs):
    r = admin.upload("/attachments", "cizim.png", _text_image("m = 2.5 mm   z = 31"), "image/png")
    doc_id = r.json()["id"]
    process_jobs()
    doc = admin.get(f"/documents/{doc_id}").json()
    assert doc["current_version"]["ingestion_status"] == "indexed", doc["current_version"]
    vid = doc["current_version"]["id"]
    pages = admin.get(f"/documents/{doc_id}/versions/{vid}/pages").json()
    assert pages[0]["extraction_method"] == "ocr"
    assert pages[0]["confidence_status"] == "draft_extraction"
    vals = admin.get(f"/extractions/values?document_id={doc_id}").json()["items"]
    labels = {v["label"]: v for v in vals}
    assert "m" in labels and labels["m"]["status"] == "draft_extraction"
    assert labels["m"]["value"] == pytest.approx(2.5)
    # preview thumbnail exists
    assert admin.get(f"/documents/{doc_id}/versions/{vid}/preview").status_code == 200


def test_pdf_page_preview_and_text(admin, process_jobs, settings):
    p = make_pdf(settings.data_root / "tmp" / "preview.pdf", ["Sayfa bir", "Sayfa iki"])
    r = admin.upload("/documents/upload", "onizleme.pdf", p.read_bytes(), "application/pdf")
    doc_id = r.json()["id"]
    process_jobs()
    vid = admin.get(f"/documents/{doc_id}").json()["current_version"]["id"]
    img = admin.get(f"/documents/{doc_id}/versions/{vid}/preview?page=2")
    assert img.status_code == 200 and img.headers["content-type"] == "image/png"
    assert admin.get(f"/documents/{doc_id}/versions/{vid}/pages/2").json()["text"].startswith("Sayfa iki")
    assert audit_rows("document.view", target_id=doc_id)


def test_document_list_filters(admin):
    r = admin.get("/documents?status=deleted")
    assert r.status_code == 200
    assert all(d["status"] == "deleted" for d in r.json()["items"])
    r2 = admin.get("/documents?category=pdf&area=ekler")
    assert all(d["current_version"]["category"] == "pdf" for d in r2.json()["items"])
