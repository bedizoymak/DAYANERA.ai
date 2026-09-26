"""Office (OOXML) and plain-text extraction."""
from __future__ import annotations

import html
import io
import re

from app.ingestion.extractors.base import ExtractContext, ExtractionOutput, PageOut, normalize_text

SEGMENT_CHARS = 3000
MAX_CHARS = 2_000_000


def _segments(text: str, label: str, method: str) -> list[PageOut]:
    text = normalize_text(text)[:MAX_CHARS]
    if not text:
        return []
    paras = re.split(r"\n\s*\n", text)
    pages, buf = [], ""
    for p in paras:
        if len(buf) + len(p) > SEGMENT_CHARS and buf:
            pages.append(buf)
            buf = ""
        buf = (buf + "\n\n" + p) if buf else p
        while len(buf) > SEGMENT_CHARS * 2:
            pages.append(buf[:SEGMENT_CHARS])
            buf = buf[SEGMENT_CHARS:]
    if buf:
        pages.append(buf)
    return [PageOut(page_number=i + 1, text=t, method=method, locator=f"{label} {i + 1}") for i, t in enumerate(pages)]


def extract_docx(data: bytes, ctx: ExtractContext) -> ExtractionOutput:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t_idx, table in enumerate(d.tables):
        rows = []
        for row in table.rows:
            rows.append(" | ".join(c.text.strip() for c in row.cells))
        parts.append(f"[Tablo {t_idx + 1}]\n" + "\n".join(rows))
    out = ExtractionOutput(metadata={"paragraphs": len(d.paragraphs), "tables": len(d.tables)})
    out.pages = _segments("\n\n".join(parts), "bölüm", "office")
    if not out.pages:
        out.stored_only, out.stored_only_reason = True, "Belgede metin bulunamadı."
    return out


def extract_xlsx(data: bytes, ctx: ExtractContext) -> ExtractionOutput:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = ExtractionOutput(metadata={"sheets": wb.sheetnames})
    n = 0
    for ws in wb.worksheets:
        lines = []
        for r_idx, row in enumerate(ws.iter_rows(values_only=True)):
            if r_idx >= 5000:
                out.warnings.append(f"'{ws.title}' sayfası 5000 satırda kesildi.")
                break
            vals = ["" if v is None else str(v) for v in row]
            if any(vals):
                lines.append(" | ".join(vals).rstrip(" |"))
        if lines:
            n += 1
            out.pages.append(PageOut(page_number=n, text=normalize_text("\n".join(lines))[:200000], method="office",
                                     locator=f"çalışma sayfası '{ws.title}'"))
    if not out.pages:
        out.stored_only, out.stored_only_reason = True, "Çalışma kitabında veri bulunamadı."
    return out


def extract_pptx(data: bytes, ctx: ExtractContext) -> ExtractionOutput:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    out = ExtractionOutput(metadata={"slides": len(prs.slides)})
    for i, slide in enumerate(prs.slides):
        texts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                texts.append(shape.text_frame.text)
        txt = normalize_text("\n".join(texts))
        if txt:
            out.pages.append(PageOut(page_number=i + 1, text=txt, method="office", locator=f"slayt {i + 1}"))
    if not out.pages:
        out.stored_only, out.stored_only_reason = True, "Sunumda metin bulunamadı."
    return out


class _HTMLStripper:
    TAGS = re.compile(r"<(script|style)[^>]*>.*?</\1>", re.S | re.I)
    ANY = re.compile(r"<[^>]+>")

    @classmethod
    def strip(cls, s: str) -> str:
        s = cls.TAGS.sub(" ", s)
        s = re.sub(r"<br\s*/?>|</p>|</div>|</li>|</h\d>", "\n", s, flags=re.I)
        return html.unescape(cls.ANY.sub(" ", s))


def decode_text(data: bytes) -> str:
    """Decode text files: UTF-16 (BOM), UTF-8, then Turkish Windows-1254."""
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        try:
            return data.decode("utf-16")
        except UnicodeDecodeError:
            pass
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass
    try:
        return data.decode("cp1254")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def extract_text(data: bytes, ctx: ExtractContext) -> ExtractionOutput:
    text = decode_text(data[:MAX_CHARS * 2])
    if ctx.filename.lower().endswith((".html", ".htm", ".xml")):
        text = _HTMLStripper.strip(text)
    out = ExtractionOutput(metadata={"characters": len(text)})
    out.pages = _segments(text, "bölüm", "plain_text")
    if not out.pages:
        out.stored_only, out.stored_only_reason = True, "Dosya boş."
    return out
